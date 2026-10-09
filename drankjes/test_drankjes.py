import json
import tempfile
import threading
import unittest
import urllib.request
from pathlib import Path

import server
from openpyxl import Workbook, load_workbook


class DrankjesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "drankjes.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_seed_order_is_lightest_first(self):
        server._write(self.path, server.seed_drinks())
        payload = server.snapshot(self.path)
        names = [drink["naam"] for drink in payload["drinks"]]
        self.assertEqual(
            names,
            [
                "Ice tea (Lipton Peach)",
                "Yugen Kombucha Passion Fruit Elderflower",
                "Almdudler",
                "Coca-Cola",
            ],
        )
        self.assertEqual(payload["drinks"][0]["suiker"], 3)
        self.assertEqual(payload["drinks"][0]["band"], "lichter")
        self.assertEqual(payload["drinks"][2]["band"], "gemiddeld")
        self.assertEqual(payload["drinks"][3]["band"], "zwaar")
        self.assertEqual(payload["drinks"][0]["rang"], 1)
        self.assertFalse(payload["warnings"])

    def test_spreadsheet_has_rank_formula_and_chart(self):
        server._write(self.path, server.seed_drinks())
        wb = load_workbook(self.path)
        ws = wb["Drankjes"]
        self.assertTrue(ws["F2"].value.startswith("=IF"))
        self.assertEqual(len(ws._charts), 1)
        self.assertTrue(ws.column_dimensions["G"].hidden)
        self.assertIn("Uitleg", wb.sheetnames)
        self.assertEqual(ws["A2"].value, "Ice tea (Lipton Peach)")

    def test_add_puts_light_drink_on_top_and_keeps_file(self):
        payload = server.add_drink(
            {"naam": "Spuitwater", "suiker": "0", "kcal": "0", "notitie": "Puur"},
            self.path,
        )
        self.assertEqual(payload["drinks"][0]["naam"], "Spuitwater")
        self.assertEqual(payload["savedId"], payload["drinks"][0]["id"])
        again = server.snapshot(self.path)
        self.assertEqual(again["drinks"][0]["naam"], "Spuitwater")
        self.assertEqual(len(again["drinks"]), 5)

    def test_comma_decimal_and_estimated_kcal(self):
        payload = server.add_drink({"naam": "Limonade", "suiker": "6,5"}, self.path)
        match = next(drink for drink in payload["drinks"] if drink["naam"] == "Limonade")
        self.assertEqual(match["suiker"], 6.5)
        self.assertEqual(match["kcal"], 26)

    def test_edit_and_delete(self):
        created = server.add_drink({"naam": "Tijdelijk", "suiker": 1, "kcal": 4}, self.path)
        drink_id = created["savedId"]
        edited = server.update_drink(
            drink_id, {"naam": "Tijdelijk light", "suiker": "0,5", "notitie": "aangepast"}, self.path
        )
        match = next(drink for drink in edited["drinks"] if drink["id"] == drink_id)
        self.assertEqual(match["naam"], "Tijdelijk light")
        self.assertEqual(match["suiker"], 0.5)
        self.assertEqual(match["notitie"], "aangepast")
        removed = server.delete_drink(drink_id, self.path)
        self.assertFalse(any(drink["id"] == drink_id for drink in removed["drinks"]))

    def test_rejects_missing_sugar_and_negative(self):
        with self.assertRaises(server.DrinkError):
            server.add_drink({"naam": "Zonder"}, self.path)
        with self.assertRaises(server.DrinkError):
            server.add_drink({"naam": "Onder nul", "suiker": -1}, self.path)

    def test_formula_text_stays_text(self):
        server.add_drink({"naam": "=1+1", "suiker": 1, "notitie": "@cmd"}, self.path)
        wb = load_workbook(self.path)
        ws = wb["Drankjes"]
        names = [ws.cell(row, 1).value for row in range(2, 7)]
        self.assertIn("=1+1", names)
        cell = next(ws.cell(row, 1) for row in range(2, 7) if ws.cell(row, 1).value == "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_excel_row_without_id_is_picked_up(self):
        server._write(self.path, server.seed_drinks())
        wb = load_workbook(self.path)
        ws = wb["Drankjes"]
        ws["A6"] = "Bronwater"
        ws["B6"] = 0
        wb.save(self.path)
        payload = server.snapshot(self.path)
        self.assertEqual(payload["drinks"][0]["naam"], "Bronwater")
        self.assertTrue(payload["drinks"][0]["id"])
        again = server.snapshot(self.path)
        self.assertEqual(again["drinks"][0]["id"], payload["drinks"][0]["id"])
        self.assertEqual(again["drinks"][0]["kcal"], 0)

    def test_invalid_row_warns_and_stays_out_of_ranking(self):
        server._write(self.path, server.seed_drinks())
        wb = load_workbook(self.path)
        ws = wb["Drankjes"]
        ws["A6"] = "Half ingevuld"
        wb.save(self.path)
        payload = server.snapshot(self.path)
        self.assertEqual(len(payload["drinks"]), 4)
        self.assertTrue(any("Half ingevuld" in warning for warning in payload["warnings"]))
        wb = load_workbook(self.path)
        self.assertEqual(wb["Drankjes"]["A6"].value, "Half ingevuld")

    def test_tie_breaks_on_kcal_then_name(self):
        server._write(
            self.path,
            [
                server.Drink("a", "Banaan", 4, 20, "", "2026-10-09"),
                server.Drink("b", "Appel", 4, 20, "", "2026-10-09"),
                server.Drink("c", "Zwaarder", 4, 30, "", "2026-10-09"),
            ],
        )
        names = [drink["naam"] for drink in server.snapshot(self.path)["drinks"]]
        self.assertEqual(names, ["Appel", "Banaan", "Zwaarder"])

    def test_glass_and_sugar_cubes(self):
        server._write(self.path, [server.Drink("a", "Test", 8, 33, "", "2026-10-09")])
        drink = server.snapshot(self.path)["drinks"][0]
        self.assertEqual(drink["glasSuiker"], 20)
        self.assertEqual(drink["klontjes"], 5)

    def test_http_add_and_page(self):
        old = server.DATA_PATH
        server.DATA_PATH = self.path
        httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        port = httpd.server_address[1]
        try:
            page = urllib.request.urlopen(f"http://127.0.0.1:{port}/").read().decode()
            self.assertIn("Drankjes die je lekker vindt", page)
            body = json.dumps({"naam": "Thee", "suiker": "0,2"}).encode()
            request = urllib.request.Request(
                f"http://127.0.0.1:{port}/api/drinks",
                data=body,
                headers={"Content-Type": "application/json"},
            )
            added = json.load(urllib.request.urlopen(request))
            self.assertEqual(added["drinks"][0]["naam"], "Thee")
            listed = json.load(urllib.request.urlopen(f"http://127.0.0.1:{port}/api/drinks"))
            self.assertEqual(listed["drinks"][0]["naam"], "Thee")
        finally:
            httpd.shutdown()
            httpd.server_close()
            thread.join(timeout=3)
            server.DATA_PATH = old

    def test_corrupt_file(self):
        self.path.write_text("geen excel", encoding="utf-8")
        with self.assertRaises(server.StoreError):
            server.snapshot(self.path)

    def test_missing_headers(self):
        wb = Workbook()
        wb.active.title = "Drankjes"
        wb.active["A1"] = "Iets anders"
        wb.save(self.path)
        with self.assertRaises(server.StoreError):
            server.snapshot(self.path)


if __name__ == "__main__":
    unittest.main()
